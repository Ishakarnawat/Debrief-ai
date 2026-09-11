import unittest
import io
import docx
from services.document_parser import DocumentParser


def create_mock_pdf(items):
    """Generates valid minimal PDF 1.4 bytes with absolute positioned text elements."""
    objects = []
    objects.append(b'<</Type /Catalog /Pages 2 0 R>>')
    objects.append(b'<</Type /Pages /Kids [3 0 R] /Count 1>>')
    objects.append(b'<</Type /Page /Parent 2 0 R /Resources <</Font <</F1 4 0 R>>>> /MediaBox [0 0 612 792] /Contents 5 0 R>>')
    objects.append(b'<</Type /Font /Subtype /Type1 /BaseFont /Helvetica>>')

    stream_ops = ['BT', '/F1 12 Tf']
    for text, x, y in items:
        # Escape parenthesis in PDF text string
        safe_text = text.replace("(", "\\(").replace(")", "\\)")
        stream_ops.append(f'1 0 0 1 {x} {y} Tm ({safe_text}) Tj')
    stream_ops.append('ET')

    stream_data = '\n'.join(stream_ops).encode('latin-1')
    stream_obj = f'<</Length {len(stream_data)}>>\nstream\n'.encode('latin-1') + stream_data + b'\nendstream'
    objects.append(stream_obj)

    out = bytearray(b'%PDF-1.4\n')
    offsets = [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out.extend(f'{i} 0 obj\n'.encode('latin-1'))
        out.extend(obj)
        out.extend(b'\nendobj\n')

    xref_offset = len(out)
    out.extend(f'xref\n0 {len(objects)+1}\n'.encode('latin-1'))
    out.extend(b'0000000000 65535 f \n')
    for off in offsets[1:]:
        out.extend(f'{off:010d} 00000 n \n'.encode('latin-1'))
    out.extend(f'trailer\n<</Size {len(objects)+1} /Root 1 0 R>>\nstartxref\n{xref_offset}\n%%EOF'.encode('latin-1'))
    return bytes(out)


class TestDocumentParser(unittest.TestCase):
    def test_parse_txt_document(self):
        txt_content = (
            "David Miller\n"
            "david.miller@example.com | +1-555-432-1098\n"
            "Summary: Senior Backend Engineer with 7 years of Python and Go experience.\n"
            "Skills: Python, FastAPI, Docker, PostgreSQL."
        ).encode("utf-8")

        parsed = DocumentParser.parse_document(txt_content, "David_Miller_Resume.txt")
        self.assertEqual(parsed.file_type, "txt")
        self.assertEqual(parsed.candidate_name, "David Miller")
        self.assertEqual(parsed.candidate_email, "david.miller@example.com")
        self.assertEqual(parsed.candidate_phone, "+1-555-432-1098")
        self.assertGreater(parsed.word_count, 15)
        self.assertIn("Senior Backend Engineer", parsed.extracted_text)

    def test_parse_docx_document(self):
        doc = docx.Document()
        doc.add_heading("Elena Rostova", level=1)
        doc.add_paragraph("elena.rostova@cloudtech.io | +1-415-555-2671")
        doc.add_paragraph("Summary: Cloud Infrastructure Architect specialized in Kubernetes and Terraform.")
        
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Certification"
        table.cell(0, 1).text = "Year"
        table.cell(1, 0).text = "AWS Solutions Architect"
        table.cell(1, 1).text = "2024"

        buf = io.BytesIO()
        doc.save(buf)
        docx_bytes = buf.getvalue()

        parsed = DocumentParser.parse_document(docx_bytes, "Elena_Rostova_Resume.docx")
        self.assertEqual(parsed.file_type, "docx")
        self.assertEqual(parsed.candidate_name, "Elena Rostova")
        self.assertEqual(parsed.candidate_email, "elena.rostova@cloudtech.io")
        self.assertEqual(parsed.candidate_phone, "+1-415-555-2671")
        self.assertIn("Cloud Infrastructure Architect", parsed.extracted_text)
        self.assertIn("AWS Solutions Architect | 2024", parsed.extracted_text)

    def test_parse_pdf_single_column(self):
        items = [
            ("Marcus Aurelius", 72, 720),
            ("marcus.aurelius@rome.org", 72, 700),
            ("+1-202-555-0143", 72, 680),
            ("Summary: Experienced Lead Architect with deep systems knowledge.", 72, 650),
            ("Skills: C++, Rust, Distributed Systems, High Availability.", 72, 630),
        ]
        pdf_bytes = create_mock_pdf(items)

        parsed = DocumentParser.parse_document(pdf_bytes, "Marcus_Aurelius_Resume.pdf")
        self.assertEqual(parsed.file_type, "pdf")
        self.assertEqual(parsed.candidate_name, "Marcus Aurelius")
        self.assertEqual(parsed.candidate_email, "marcus.aurelius@rome.org")
        self.assertEqual(parsed.candidate_phone, "+1-202-555-0143")
        self.assertIn("Experienced Lead Architect", parsed.extracted_text)
        self.assertEqual(parsed.page_count, 1)

    def test_multi_column_flow_separation(self):
        # Construct a 2-column page:
        # Width: 612. Gutter: around 250-280.
        # Left column (x=50 to 220): Contact and Skills
        # Right column (x=300 to 550): Experience items
        items = [
            ("Alice Smith", 220, 750),  # Header at top
            ("alice@example.com", 220, 730),
            # Left column words
            ("LeftColHeader: Skills", 50, 680),
            ("Python Programming", 50, 650),
            ("FastAPI Framework", 50, 620),
            ("PostgreSQL Database", 50, 590),
            ("Docker Containers", 50, 560),
            ("Git Version Control", 50, 530),
            ("Unit Testing Pytest", 50, 500),
            ("Linux Administration", 50, 470),
            # Right column words
            ("RightColHeader: Work Experience", 320, 680),
            ("Senior Engineer at Tech Corp", 320, 650),
            ("Built scalable microservices", 320, 620),
            ("Optimized database queries", 320, 590),
            ("Managed Kubernetes cluster", 320, 560),
            ("Mentored junior developers", 320, 530),
            ("Delivered 99.9% uptime", 320, 500),
            ("Led sprint retrospectives", 320, 470),
        ]
        pdf_bytes = create_mock_pdf(items)

        parsed = DocumentParser.parse_document(pdf_bytes, "Alice_Smith_Resume.pdf")
        self.assertEqual(parsed.candidate_name, "Alice Smith")
        self.assertEqual(parsed.candidate_email, "alice@example.com")
        self.assertIn("Python Programming", parsed.extracted_text)
        self.assertIn("Work Experience", parsed.extracted_text)

    def test_metadata_extractor_heuristics(self):
        sample = """
        Johnathan Edward Doe
        contact@johndoe.dev | (555) 345-6789 | San Francisco, CA
        https://linkedin.com/in/johndoe

        Professional Experience
        Lead Software Architect, Google (2020 - Present)
        """
        name = DocumentParser.extract_name(sample, "Resume_2026.pdf")
        email = DocumentParser.extract_email(sample)
        phone = DocumentParser.extract_phone(sample)

        self.assertEqual(name, "Johnathan Edward Doe")
        self.assertEqual(email, "contact@johndoe.dev")
        self.assertIn("555", phone)


if __name__ == "__main__":
    unittest.main()
